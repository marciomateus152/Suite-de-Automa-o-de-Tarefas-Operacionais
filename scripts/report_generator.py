#!/usr/bin/env python3
"""
report_generator.py - Gerador de Relatórios CSV/Excel
Lê múltiplos CSVs, realiza merge, limpeza e exporta relatório formatado.
"""

import logging
import sys
import time
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# Configuração de logging
# ---------------------------------------------------------------------------

LOG_DIR = Path(__file__).parent.parent / "logs"
LOG_DIR.mkdir(exist_ok=True)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s | %(levelname)-8s | %(message)s",
    datefmt="%Y-%m-%d %H:%M:%S",
    handlers=[
        logging.FileHandler(LOG_DIR / "report_generator.log", encoding="utf-8"),
        logging.StreamHandler(sys.stdout),
    ],
)
logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Constantes de formatação Excel
# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
ALT_ROW_FILL = PatternFill(start_color="D6E4F0", end_color="D6E4F0", fill_type="solid")
DEFAULT_FONT = Font(name="Calibri", size=10)

# ---------------------------------------------------------------------------
# Leitura e merge de CSVs
# ---------------------------------------------------------------------------


def load_csvs(csv_dir: Path, encoding: str = "utf-8") -> pd.DataFrame:
    """
    Lê todos os arquivos CSV de um diretório e retorna o DataFrame combinado.

    Args:
        csv_dir: Diretório contendo os arquivos CSV.
        encoding: Codificação dos arquivos (padrão utf-8).

    Returns:
        DataFrame com todos os dados combinados.
    """
    csv_files = sorted(csv_dir.glob("*.csv"))
    if not csv_files:
        raise FileNotFoundError(f"Nenhum CSV encontrado em: {csv_dir}")

    frames: list[pd.DataFrame] = []
    for csv_path in csv_files:
        try:
            df = pd.read_csv(csv_path, encoding=encoding, dtype=str)
            df["_source_file"] = csv_path.name
            frames.append(df)
            logger.info("CSV carregado: '%s' (%d linhas)", csv_path.name, len(df))
        except (OSError, pd.errors.ParserError) as exc:
            logger.warning("Erro ao ler '%s': %s — arquivo ignorado.", csv_path.name, exc)

    if not frames:
        raise ValueError("Nenhum CSV foi carregado com sucesso.")

    combined = pd.concat(frames, ignore_index=True)
    logger.info("Total combinado antes da limpeza: %d linhas.", len(combined))
    return combined


# ---------------------------------------------------------------------------
# Limpeza de dados
# ---------------------------------------------------------------------------


def clean_dataframe(df: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """
    Executa pipeline de limpeza: remove colunas vazias, linhas nulas e duplicatas.

    Args:
        df: DataFrame bruto.

    Returns:
        Tupla (DataFrame limpo, métricas de limpeza).
    """
    metrics = {"initial_rows": len(df), "initial_cols": len(df.columns)}

    # Remove colunas 100% vazias
    empty_cols = [c for c in df.columns if df[c].isna().all()]
    if empty_cols:
        df = df.drop(columns=empty_cols)
        logger.info("Colunas vazias removidas: %s", empty_cols)
    metrics["empty_cols_removed"] = len(empty_cols)

    # Remove linhas completamente nulas
    before = len(df)
    df = df.dropna(how="all")
    metrics["blank_rows_removed"] = before - len(df)

    # Normaliza espaços em strings
    str_cols = df.select_dtypes(include="object").columns
    df[str_cols] = df[str_cols].apply(lambda s: s.str.strip())

    # Remove duplicatas (exclui a coluna de origem do critério)
    dupe_cols = [c for c in df.columns if c != "_source_file"]
    before = len(df)
    df = df.drop_duplicates(subset=dupe_cols, keep="first")
    metrics["duplicates_removed"] = before - len(df)
    logger.info("Duplicatas removidas: %d", metrics["duplicates_removed"])

    # Remove a coluna auxiliar de rastreio de origem
    df = df.drop(columns=["_source_file"], errors="ignore")

    metrics["final_rows"] = len(df)
    metrics["final_cols"] = len(df.columns)
    return df, metrics


# ---------------------------------------------------------------------------
# Exportação para CSV
# ---------------------------------------------------------------------------


def export_csv(df: pd.DataFrame, output_path: Path) -> None:
    """Exporta o DataFrame para CSV com separador ponto-e-vírgula."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(output_path, index=False, sep=";", encoding="utf-8-sig")
    logger.info("CSV exportado: %s (%d linhas)", output_path, len(df))


# ---------------------------------------------------------------------------
# Exportação para Excel formatado
# ---------------------------------------------------------------------------


def export_excel(df: pd.DataFrame, output_path: Path) -> None:
    """
    Exporta o DataFrame para Excel com formatação profissional:
    cabeçalho colorido, linhas alternadas, ajuste automático de colunas e filtros.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)
    df.to_excel(output_path, index=False, sheet_name="Relatório")

    wb = load_workbook(output_path)
    ws = wb.active

    # Cabeçalho
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Linhas de dados com zebra striping
    for row_idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        fill = ALT_ROW_FILL if row_idx % 2 == 0 else None
        for cell in row:
            cell.font = DEFAULT_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center")
            if fill:
                cell.fill = fill

    # Ajuste automático de largura de colunas
    for col_idx, col_cells in enumerate(ws.columns, start=1):
        max_len = max((len(str(c.value or "")) for c in col_cells), default=10)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 4, 50)

    # Filtros automáticos na linha de cabeçalho
    ws.auto_filter.ref = ws.dimensions

    # Congelar painel no cabeçalho
    ws.freeze_panes = "A2"

    wb.save(output_path)
    logger.info("Excel formatado exportado: %s", output_path)


# ---------------------------------------------------------------------------
# Relatório de métricas
# ---------------------------------------------------------------------------


def print_metrics(metrics: dict, elapsed: float, csv_out: Path, xlsx_out: Path) -> None:
    sep = "=" * 55
    print(f"\n{sep}")
    print("  RELATÓRIO DE GERAÇÃO — MÉTRICAS")
    print(sep)
    print(f"  Linhas iniciais          : {metrics['initial_rows']}")
    print(f"  Colunas vazias removidas : {metrics['empty_cols_removed']}")
    print(f"  Linhas em branco         : {metrics['blank_rows_removed']}")
    print(f"  Duplicatas removidas     : {metrics['duplicates_removed']}")
    print(f"  Linhas no relatório final: {metrics['final_rows']}")
    print(f"  Colunas finais           : {metrics['final_cols']}")
    print(f"  Tempo de processamento   : {elapsed:.3f}s")
    print(f"\n  Arquivos gerados:")
    print(f"    CSV  → {csv_out}")
    print(f"    XLSX → {xlsx_out}")
    print(sep)


# ---------------------------------------------------------------------------
# Ponto de entrada
# ---------------------------------------------------------------------------


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser(
        description="Gera relatório consolidado a partir de múltiplos CSVs."
    )
    base = Path(__file__).parent.parent
    parser.add_argument(
        "--input-dir",
        default=str(base / "input" / "sample_csvs"),
        help="Diretório com os arquivos CSV de entrada",
    )
    parser.add_argument(
        "--output-dir",
        default=str(base / "output"),
        help="Diretório de saída para os relatórios",
    )
    parser.add_argument(
        "--report-name",
        default="relatorio_consolidado",
        help="Nome base dos arquivos de saída (sem extensão)",
    )
    parser.add_argument(
        "--encoding",
        default="utf-8",
        help="Codificação dos CSVs de entrada",
    )
    args = parser.parse_args()

    input_dir = Path(args.input_dir)
    output_dir = Path(args.output_dir)
    start = time.perf_counter()

    try:
        df_raw = load_csvs(input_dir, encoding=args.encoding)
        df_clean, metrics = clean_dataframe(df_raw)

        csv_out = output_dir / f"{args.report_name}.csv"
        xlsx_out = output_dir / f"{args.report_name}.xlsx"

        export_csv(df_clean, csv_out)
        export_excel(df_clean, xlsx_out)

        elapsed = time.perf_counter() - start
        print_metrics(metrics, elapsed, csv_out, xlsx_out)
        logger.info("Pipeline concluído em %.3fs.", elapsed)

    except (FileNotFoundError, ValueError) as exc:
        logger.critical("Erro crítico: %s", exc)
        sys.exit(1)


if __name__ == "__main__":
    main()
